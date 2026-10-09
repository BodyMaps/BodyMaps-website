"""Read existing PanTS assets only; no production API, database or job access."""
from __future__ import annotations

import hashlib
import hmac
import json
import math
import os
import re
import threading
import time
from collections import defaultdict, deque
from pathlib import Path

from flask import Flask, abort, g, jsonify, request, send_file
from openpyxl import load_workbook

CASE_RE = re.compile(r"(?:PanTS_)?([0-9]{1,8})\Z", re.ASCII)
MESH_RE = re.compile(r"[a-zA-Z0-9_-]{1,120}\.glb\Z", re.ASCII)
CASE_COLUMNS = ("PanTS ID", "PanTS_ID", "case_id", "id", "case", "CaseID")


def case_id(value):
    match = CASE_RE.fullmatch(str(value))
    if not match or int(match[1]) < 1:
        raise ValueError("Invalid PanTS case ID")
    return f"PanTS_{int(match[1]):08d}"


def existing_file(root: Path, relative: str) -> Path:
    """Allow no file outside a configured export root, including symlink escapes."""
    path = (root / relative).resolve(strict=True)
    path.relative_to(root.resolve(strict=True))
    if not path.is_file():
        raise FileNotFoundError(relative)
    return path


def json_value(value):
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def read_metadata(root):
    try:
        path = existing_file(root, "metadata.xlsx")
    except FileNotFoundError:
        path = existing_file(root, "data/metadata.xlsx")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook["PanTS_metadata"] if "PanTS_metadata" in workbook.sheetnames else workbook.active
        rows = sheet.iter_rows(values_only=True)
        columns = [str(v or "") for v in next(rows)]
        index = next((columns.index(c) for c in CASE_COLUMNS if c in columns), None)
        if index is None:
            raise ValueError("No recognized case ID column in metadata")
        catalog = {}
        for row in rows:
            try:
                identifier = case_id(row[index])
            except ValueError:
                continue
            catalog[identifier] = [json_value(v) for v in row]
        return columns, catalog
    finally:
        workbook.close()


def create_app(config=None):
    app = Flask(__name__, static_folder=None)
    app.config.update(
        DATA_ROOT=os.environ.get("BODYMAPS_RO_DATA_ROOT", ""),
        MESH_ROOT=os.environ.get("BODYMAPS_RO_MESH_ROOT", ""),
        TOKENS_FILE=os.environ.get("BODYMAPS_RO_TOKENS_FILE", ""),
        REQUESTS_PER_MINUTE=120,
        MAX_CONTENT_LENGTH=0,
    )
    app.config.update(config or {})
    if not app.config["DATA_ROOT"] or not app.config["TOKENS_FILE"]:
        raise ValueError("DATA_ROOT and TOKENS_FILE must be configured")
    data_root = Path(app.config["DATA_ROOT"]).resolve(strict=True)
    mesh_root = Path(app.config["MESH_ROOT"]).resolve(strict=True) if app.config["MESH_ROOT"] else None
    columns, metadata = read_metadata(data_root)
    attempts = defaultdict(deque)
    lock = threading.Lock()

    def token_records():
        # Reload every request: deleting a token immediately revokes access.
        records = json.loads(Path(app.config["TOKENS_FILE"]).read_text(encoding="utf-8"))
        if not isinstance(records, list) or len(records) > 1000:
            raise ValueError("Invalid token configuration")
        for record in records:
            if not isinstance(record, dict):
                raise ValueError("Invalid token record")
            if not re.fullmatch(r"[a-f0-9]{64}", record["sha256"]):
                raise ValueError("Invalid token hash")
            if not isinstance(record["cases"], list) or not record["cases"]:
                raise ValueError("Explicit approved cases are required")
            record["cases"] = {case_id(c) for c in record["cases"]}
            expiry = record["expires_at"]
            if isinstance(expiry, bool) or not isinstance(expiry, (int, float)) or not math.isfinite(expiry):
                raise ValueError("A finite expiry timestamp is required")
        return records

    token_records()  # Fail closed at startup on invalid configuration.

    @app.before_request
    def protect():
        if request.method not in {"GET", "HEAD"}:
            abort(405)
        if request.content_length or request.headers.get("Transfer-Encoding"):
            abort(400)
        # No arbitrary query controls, URL fetching, cookies or production sessions.
        if request.query_string:
            abort(400)
        now = time.monotonic()
        with lock:
            # Service binds loopback, so this also bounds unauthenticated traffic
            # across SSH tunnels without trusting spoofable forwarded headers.
            bucket = attempts["all"]
            while bucket and now - bucket[0] >= 60:
                bucket.popleft()
            if len(bucket) >= app.config["REQUESTS_PER_MINUTE"]:
                return jsonify(error="Read limit reached; retry later"), 429, {"Retry-After": "60"}
            bucket.append(now)
        auth = request.headers.get("Authorization", "")
        if not auth.startswith("Bearer ") or not 32 <= len(auth[7:]) <= 256:
            abort(401)
        digest = hashlib.sha256(auth[7:].encode()).hexdigest()
        try:
            records = token_records()
        except (OSError, ValueError, KeyError, TypeError):
            abort(503, "Access configuration unavailable")
        principal = next((r for r in records if hmac.compare_digest(r["sha256"], digest)), None)
        if principal is None or float(principal["expires_at"]) <= time.time():
            abort(401)
        g.approved_cases = principal["cases"]

    @app.after_request
    def headers(response):
        response.headers["Cache-Control"] = "private, no-store"
        response.headers["X-Content-Type-Options"] = "nosniff"
        return response

    def authorize(identifier):
        try:
            identifier = case_id(identifier)
        except ValueError:
            abort(404)
        if identifier not in g.approved_cases or identifier not in metadata:
            abort(404)
        return identifier

    def assets(identifier):
        candidates = {
            "ct": (data_root, f"image_only/{identifier}/ct.nii.gz"),
            "mask": (data_root, f"mask_only/{identifier}/combined_labels.nii.gz"),
            "thumbnail": (data_root, f"profile_only/{identifier}/profile.jpg"),
        }
        # Only existing GLBs are exposed. Missing meshes are generated LOCALLY
        # by the normal backend after import, never on this export service.
        if mesh_root:
            try:
                manifest_path = existing_file(mesh_root, f"{identifier}/manifest.json")
                if manifest_path.stat().st_size > 1024 * 1024:
                    raise ValueError("Oversized mesh manifest")
                manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
                if not isinstance(manifest, dict) or not isinstance(manifest.get("organs"), list):
                    raise ValueError("Invalid mesh manifest")
                candidates["mesh-manifest"] = (mesh_root, f"{identifier}/manifest.json")
                for organ in manifest.get("organs", []):
                    if not isinstance(organ, dict):
                        continue
                    filename = str(organ.get("key", "")) + ".glb"
                    if MESH_RE.fullmatch(filename):
                        candidates[filename] = (mesh_root, f"{identifier}/{filename}")
            except (OSError, ValueError):
                pass
        found = {}
        for key, (root, relative) in candidates.items():
            try:
                found[key] = existing_file(root, relative)
            except (OSError, ValueError):
                pass
        return found

    @app.get("/v1/catalog")
    def catalog():
        return jsonify(schema=1, dataset="PanTS", cases=sorted(g.approved_cases & metadata.keys()))

    @app.get("/v1/cases/<identifier>")
    def describe_case(identifier):
        identifier = authorize(identifier)
        return jsonify(
            schema=1, case_id=identifier, columns=columns, row=metadata[identifier],
            assets={key: {"bytes": path.stat().st_size} for key, path in assets(identifier).items()},
        )

    @app.get("/v1/cases/<identifier>/assets/<asset>")
    def download_asset(identifier, asset):
        identifier = authorize(identifier)
        path = assets(identifier).get(asset)
        if path is None:
            abort(404)
        return send_file(path, conditional=True, etag=True, as_attachment=True)

    return app
