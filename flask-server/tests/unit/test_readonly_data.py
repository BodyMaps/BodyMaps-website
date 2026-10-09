"""Exercise the export boundary without production credentials or real scans."""
import hashlib
import json
import os
from pathlib import Path
import sys
import subprocess
import socket
import threading
import time

import pytest
from openpyxl import Workbook, load_workbook
from werkzeug.serving import make_server

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from readonly_data.client import asset_relative, fetch_snapshot, validate_server
from readonly_data.credentials import prepare_grant
from readonly_data.local import local_environment
from readonly_data.service import create_app

TOKEN = "test-only-not-a-production-token-" * 2


@pytest.fixture
def export(tmp_path):
    data = tmp_path / "source"
    meshes = tmp_path / "meshes"
    data.mkdir()
    meshes.mkdir()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "PanTS_metadata"
    sheet.append(["PanTS ID", "Sex", "Age"])
    for identifier in (35, 36):
        canonical = f"PanTS_{identifier:08d}"
        sheet.append([canonical, "F", 45])
        for relative, content in [(f"image_only/{canonical}/ct.nii.gz", b"synthetic-ct"),
                                  (f"mask_only/{canonical}/combined_labels.nii.gz", b"synthetic-mask"),
                                  (f"profile_only/{canonical}/profile.jpg", b"synthetic-thumbnail")]:
            path = data / relative
            path.parent.mkdir(parents=True)
            path.write_bytes(content)
    workbook.save(data / "metadata.xlsx")
    case = meshes / "PanTS_00000035"
    case.mkdir()
    (case / "manifest.json").write_text(json.dumps({"organs": [{"key": "liver"}, {"key": "../../secret"}]}))
    (case / "liver.glb").write_bytes(b"synthetic-glb")
    (case / "private.txt").write_text("not exported")
    records = [{"sha256": hashlib.sha256(TOKEN.encode()).hexdigest(), "cases": [35], "expires_at": time.time() + 600}]
    tokens = tmp_path / "tokens.json"
    tokens.write_text(json.dumps(records))
    config = {"TESTING": True, "DATA_ROOT": str(data), "MESH_ROOT": str(meshes), "TOKENS_FILE": str(tokens)}
    app = create_app(config)
    return app, {"Authorization": "Bearer " + TOKEN}, data, meshes, tokens, records, config


def test_scoped_catalog_and_exact_asset_allowlist(export):
    app, headers, *_ = export
    client = app.test_client()
    assert client.get("/v1/catalog", headers=headers).json["cases"] == ["PanTS_00000035"]
    description = client.get("/v1/cases/35", headers=headers).json
    assert description["row"] == ["PanTS_00000035", "F", 45]
    assert set(description["assets"]) == {"ct", "mask", "thumbnail", "mesh-manifest", "liver.glb"}
    for endpoint in ("/v1/cases/36", "/v1/cases/36/assets/ct", "/v1/cases/35/assets/private.txt",
                     "/api/jobs/next", "/api/meshes/35/manifest", "/v1/cases/35/assets/../../tokens.json",
                     "/v1/cases/%2e%2e/assets/ct", "/v1/cases/35/assets/%2fetc%2fpasswd"):
        assert client.get(endpoint, headers=headers).status_code == 404


def test_unauthorized_expired_and_revoked_tokens(export):
    app, headers, _, _, tokens, records, _ = export
    client = app.test_client()
    assert client.get("/v1/catalog").status_code == 401
    assert client.get("/v1/catalog", headers={"Authorization": "Bearer " + "x" * 64}).status_code == 401
    records[0]["expires_at"] = time.time() - 1
    tokens.write_text(json.dumps(records))
    assert client.get("/v1/catalog", headers=headers).status_code == 401
    tokens.write_text("[]")
    assert client.get("/v1/catalog", headers=headers).status_code == 401
    tokens.write_text("not valid json")
    assert client.get("/v1/catalog", headers=headers).status_code == 503


@pytest.mark.parametrize("expiry", [float("nan"), float("inf"), "tomorrow", True, None])
def test_invalid_expiry_fails_closed(export, expiry):
    _, _, _, _, tokens, records, config = export
    records[0]["expires_at"] = expiry
    tokens.write_text(json.dumps(records))
    with pytest.raises(ValueError):
        create_app(config)


@pytest.mark.parametrize("method", ["POST", "PUT", "PATCH", "DELETE", "OPTIONS"])
def test_write_methods_rejected(export, method):
    app, headers, *_ = export
    assert app.test_client().open("/v1/cases/35/assets/ct", method=method, headers=headers).status_code == 405


def test_body_query_rate_limit_and_range(export):
    app, headers, *_ = export
    client = app.test_client()
    assert client.get("/v1/catalog?url=http://localhost", headers=headers).status_code == 400
    assert client.get("/v1/catalog", headers=headers, data="payload").status_code == 400
    response = client.get("/v1/cases/35/assets/ct", headers={**headers, "Range": "bytes=0-3"})
    assert response.status_code == 206 and response.data == b"synt"
    response.close()
    response = client.head("/v1/cases/35/assets/ct", headers=headers)
    assert response.status_code == 200 and response.data == b""
    assert response.headers["Cache-Control"] == "private, no-store"
    response.close()
    app.config["REQUESTS_PER_MINUTE"] = 2
    assert client.get("/v1/catalog", headers=headers).status_code == 429


def test_missing_files_do_not_generate_and_symlinks_cannot_escape(export, tmp_path):
    app, headers, data, _, *_ = export
    source = data / "image_only/PanTS_00000035/ct.nii.gz"
    source.unlink()
    client = app.test_client()
    assert client.get("/v1/cases/35/assets/ct", headers=headers).status_code == 404
    assert not source.exists()
    outside = tmp_path / "secret"
    outside.write_text("secret")
    try:
        source.symlink_to(outside)
    except OSError:
        pytest.skip("Symlink creation needs Windows Developer Mode; exercised on Linux CI")
    assert client.get("/v1/cases/35/assets/ct", headers=headers).status_code == 404


@pytest.fixture
def running_export(export):
    server = make_server("127.0.0.1", 0, export[0])
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"http://127.0.0.1:{server.server_port}"
    finally:
        server.shutdown()
        thread.join(timeout=5)
        server.server_close()


def test_http_snapshot_and_local_environment_are_isolated(export, running_export, tmp_path):
    _, _, data, meshes, *_ = export
    source_files = [path for root in (data, meshes) for path in root.rglob("*") if path.is_file()]
    before = {path: (hashlib.sha256(path.read_bytes()).hexdigest(), path.stat().st_mtime_ns) for path in source_files}
    token = tmp_path / "client.token"
    token.write_text(TOKEN)
    snapshot = fetch_snapshot(running_export, token, [35], tmp_path / "snapshot")
    assert not (snapshot / "INCOMPLETE").exists()
    assert (snapshot / "PanTS/image_only/PanTS_00000035/ct.nii.gz").read_bytes() == b"synthetic-ct"
    workbook = load_workbook(snapshot / "PanTS/metadata.xlsx", read_only=True)
    assert list(workbook.active.values) == [("PanTS ID", "Sex", "Age"), ("PanTS_00000035", "F", 45)]
    workbook.close()
    state = tmp_path / "local-state"
    env = local_environment(snapshot, state, {"PATH": "local-path", "DATABASE_URL": "postgresql://production",
                                             "EPAI_REMOTE_HOST": "production", "SECRET_KEY": "production",
                                             "VITE_API_BASE": "https://production"})
    assert env["DATABASE_URL"] == "sqlite:///" + (state / "test.db").as_posix()
    assert "SECRET_KEY" not in env and "EPAI_REMOTE_HOST" not in env
    assert env["EPAI_REMOTE_ENABLED"] == "false" and env["PYTHON_DOTENV_DISABLED"] == "1"
    assert env["VITE_PROXY_TARGET"] == "http://127.0.0.1:5001" and env["VITE_API_BASE"] == ""
    assert env["PATH"] == "local-path"
    assert before == {path: (hashlib.sha256(path.read_bytes()).hexdigest(), path.stat().st_mtime_ns) for path in source_files}
    with pytest.raises(FileExistsError):
        fetch_snapshot(running_export, token, [35], snapshot)


def test_failed_or_oversized_snapshot_cannot_start_backend(export, running_export, tmp_path):
    token = tmp_path / "client.token"
    token.write_text(TOKEN)
    output = tmp_path / "partial"
    with pytest.raises(ValueError, match="limit"):
        fetch_snapshot(running_export, token, [35], output, max_bytes=1)
    assert (output / "INCOMPLETE").exists()
    with pytest.raises(ValueError, match="completed"):
        local_environment(output, tmp_path / "state")


@pytest.mark.parametrize("url", ["http://production", "https://user:secret@production", "https://production/api", "https://production?token=secret"])
def test_unsafe_transport_rejected(url):
    with pytest.raises(ValueError):
        validate_server(url)


@pytest.mark.parametrize("asset", ["../secret", "..\\secret", "/etc/passwd", "secret.txt", "C:secret.glb"])
def test_remote_asset_names_cannot_choose_local_paths(asset):
    with pytest.raises(ValueError):
        asset_relative("35", asset)


def test_generated_grant_contains_only_hash_and_explicit_scope(tmp_path):
    token_file, grant_file = tmp_path / "secret.token", tmp_path / "grant.json"
    grant = prepare_grant([35, "PanTS_00000035"], 1, token_file, grant_file)
    token = token_file.read_text().strip()
    assert grant["sha256"] == hashlib.sha256(token.encode()).hexdigest()
    assert token not in grant_file.read_text() and grant["cases"] == ["PanTS_00000035"]
    with pytest.raises(ValueError):
        prepare_grant([36], 1, token_file, grant_file)


def test_snapshot_does_not_follow_redirects(export, running_export, tmp_path):
    from flask import redirect
    app = export[0]
    app.view_functions["describe_case"] = lambda **kwargs: redirect("https://example.invalid/")
    token = tmp_path / "client.token"
    token.write_text(TOKEN)
    with pytest.raises(RuntimeError, match="302"):
        fetch_snapshot(running_export, token, [35], tmp_path / "snapshot")
    assert (tmp_path / "snapshot/INCOMPLETE").exists()


def test_local_launcher_disables_dotenv_in_a_fresh_process(tmp_path):
    pytest.importorskip("dotenv")
    snapshot = tmp_path / "snapshot"
    (snapshot / "PanTS").mkdir(parents=True)
    (snapshot / "PanTS/metadata.xlsx").touch()
    (snapshot / "snapshot.json").write_text('{"schema": 1}')
    env_file = tmp_path / ".env"
    env_file.write_text("PRODUCTION_TEST_CREDENTIAL=must-not-load\nDATABASE_URL=postgresql://production\n")
    env = local_environment(snapshot, tmp_path / "state")
    check = "import os,sys; from dotenv import load_dotenv; load_dotenv(sys.argv[1],override=True); assert 'PRODUCTION_TEST_CREDENTIAL' not in os.environ; assert os.environ['DATABASE_URL'].startswith('sqlite:///')"
    subprocess.run([sys.executable, "-c", check, str(env_file)], env=env, check=True)


@pytest.mark.skipif(os.name == "nt", reason="Gunicorn is Linux/Unix only; exercised by CI")
def test_gunicorn_entrypoint_is_independent_of_production(export, tmp_path):
    import requests
    _, headers, data, meshes, tokens, *_ = export
    with socket.socket() as probe:
        probe.bind(("127.0.0.1", 0))
        port = probe.getsockname()[1]
    env = {key: os.environ[key] for key in ("PATH", "HOME", "LANG") if key in os.environ}
    env.update({"BODYMAPS_RO_DATA_ROOT": str(data), "BODYMAPS_RO_MESH_ROOT": str(meshes),
                "BODYMAPS_RO_TOKENS_FILE": str(tokens), "PYTHONDONTWRITEBYTECODE": "1",
                "PYTHON_DOTENV_DISABLED": "1"})
    log_path = tmp_path / "gunicorn.log"
    with log_path.open("w") as log:
        process = subprocess.Popen([sys.executable, "-m", "gunicorn", "--workers", "1", "--threads", "2",
                                    "--bind", f"127.0.0.1:{port}", "readonly_data.service:create_app()"],
                                   cwd=Path(__file__).resolve().parents[2], env=env, stdout=log, stderr=log)
        try:
            with requests.Session() as session:
                session.trust_env = False
                for _ in range(50):
                    if process.poll() is not None:
                        pytest.fail(log_path.read_text())
                    try:
                        response = session.get(f"http://127.0.0.1:{port}/v1/catalog", headers=headers, timeout=1)
                        assert response.status_code == 200
                        assert response.json()["cases"] == ["PanTS_00000035"]
                        break
                    except requests.ConnectionError:
                        time.sleep(0.1)
                else:
                    pytest.fail("Gunicorn did not become ready")
        finally:
            process.terminate()
            try:
                process.wait(timeout=10)
            except subprocess.TimeoutExpired:
                process.kill()
                process.wait(timeout=5)
